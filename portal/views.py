from decimal import Decimal, InvalidOperation
from django.utils import timezone

from django.contrib import messages
from django.http import HttpResponseNotAllowed
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.http import require_POST
from openpyxl import load_workbook

from . import auth, supabase


def _normalize_header(value):
    return ''.join(character.lower() for character in str(value or '') if character.isalnum())


def _find_column(headers, patterns):
    return next((index for index, header in enumerate(headers) if any(pattern in header for pattern in patterns)), None)


def _cell(row, index):
    if index is None or index >= len(row) or row[index] is None:
        return ''
    return str(row[index]).strip()


def _api_error(request, error):
    messages.error(request, str(error))


def _list(request, table, **params):
    return supabase.rest_list(table, params, token=request.supabase_token)


def _one(request, table, **params):
    rows = _list(request, table, **params)
    return rows[0] if rows else None


def _student_for_profile(request):
    profile = request.profile
    return _one(request, 'students', profile_id=f"eq.{profile['id']}")


def home(request):
    if request.session.get('profile'):
        return redirect('dashboard')
    return redirect('login')


def login_view(request):
    if request.session.get('profile') and auth.get_access_token(request):
        return redirect('dashboard')
    if request.method == 'POST':
        admission_number = request.POST.get('admission_number', '').strip()
        password = request.POST.get('password', '')
        if not admission_number or not password:
            messages.error(request, 'Enter your admission number and password.')
        else:
            try:
                auth.sign_in(request, admission_number, password)
                return redirect('dashboard')
            except supabase.SupabaseError as error:
                messages.error(request, str(error))
    elif request.method != 'GET':
        return HttpResponseNotAllowed(['GET', 'POST'])
    return render(request, 'portal/login.html')


def logout_view(request):
    request.session.flush()
    return redirect('login')


@auth.role_required('admin', 'student', 'principal')
def dashboard(request):
    role = request.profile['role']
    if role == 'admin':
        return redirect('admin-section', section='')
    if role == 'principal':
        return redirect('principal-section', section='home')
    return redirect('student-section', section='home')


@auth.role_required('admin')
def admin_section(request, section):
    section = section or 'home'
    section_titles = {
        'profiles': 'Profiles', 'students': 'Students', 'schedule': 'Schedule',
        'activities': 'Activities', 'attendance': 'Attendance',
        'attendance-requests': 'Attendance Requests', 'student-reports': 'Student Reports',
        'funds': 'Trip Fund', 'photos': 'Photos', 'notifications': 'Notifications',
        'settings': 'Settings', 'home': 'Dashboard', 'import': 'Import Students',
    }
    title = section_titles.get(section)
    if not title:
        return redirect('admin-section', section='home')
    data = {'rows': [], 'columns': [], 'summary': {}, 'section': section}
    try:
        if section == 'home':
            trip = _one(request, 'trips')
            students = _list(request, 'students')
            payments = _list(request, 'payments')
            activities = _list(request, 'activities', order='start_time.asc')
            total_expected = sum((Decimal(str(student['trip_fee'])) for student in students), Decimal('0'))
            total_collected = sum((Decimal(str(payment['amount'])) for payment in payments), Decimal('0'))
            data.update({
                'trip': trip,
                'summary': {
                    'Students': len(students),
                    'Expected': total_expected,
                    'Collected': total_collected,
                    'Balance': total_expected - total_collected,
                },
                'rows': activities,
                'columns': [('title', 'Activity'), ('start_time', 'Starts'), ('status', 'Status')],
            })
        elif section == 'profiles':
            data.update(rows=_list(request, 'profiles', order='created_at.desc'), columns=[('full_name', 'Name'), ('role', 'Role'), ('phone', 'Phone')])
        elif section == 'students':
            students = _list(request, 'students', order='serial_number.asc')
            payments = _list(request, 'payments')
            paid = {}
            for payment in payments:
                paid[payment['student_id']] = paid.get(payment['student_id'], Decimal('0')) + Decimal(str(payment['amount']))
            for student in students:
                student['paid_total'] = paid.get(student['id'], Decimal('0'))
                student['balance'] = Decimal(str(student['trip_fee'])) - student['paid_total']
                student['href'] = f"/admin/students/{student['id']}/"
            data.update(rows=students, columns=[('serial_number', '#'), ('full_name', 'Student'), ('admission_number', 'Login ID'), ('trip_fee', 'Fee'), ('paid_total', 'Paid'), ('balance', 'Balance')], action_url='/admin/students/new/', action_label='Add Student')
        elif section == 'schedule':
            days = _list(request, 'trip_days', order='day_number.asc')
            activities = _list(request, 'activities', order='sort_order.asc')
            data.update(rows=[{**day, 'activities': [activity for activity in activities if activity['trip_day_id'] == day['id']]} for day in days], columns=[('day_number', 'Day'), ('trip_date', 'Date'), ('title', 'Itinerary')])
        elif section == 'activities':
            days = _list(request, 'trip_days', order='day_number.asc')
            data.update(rows=_list(request, 'activities', order='start_time.asc'), columns=[('title', 'Activity'), ('location', 'Location'), ('start_time', 'Start'), ('status', 'Status')], options={'trip_days': days})
        elif section == 'attendance':
            sessions = _list(request, 'attendance_sessions', order='started_at.desc')
            activities = {item['id']: item for item in _list(request, 'activities')}
            records = _list(request, 'attendance_records')
            for session in sessions:
                activity = activities.get(session['activity_id'], {})
                session['activity_title'] = activity.get('title', 'Activity')
                session['present_count'] = sum(1 for record in records if record['session_id'] == session['id'] and record['status'] == 'present')
            data.update(rows=sessions, columns=[('activity_title', 'Activity'), ('started_at', 'Started'), ('status', 'Status'), ('present_count', 'Present')], options={'activities': [item for item in activities.values() if item.get('status') == 'live']})
        elif section == 'attendance-requests':
            requests = _list(request, 'attendance_requests', order='created_at.desc')
            students = {item['id']: item for item in _list(request, 'students')}
            activities = {item['id']: item for item in _list(request, 'activities')}
            sessions = {item['id']: item for item in _list(request, 'attendance_sessions')}
            for item in requests:
                item['student_name'] = students.get(item['student_id'], {}).get('full_name', 'Student')
                session = sessions.get(item['session_id'], {})
                item['activity_title'] = activities.get(session.get('activity_id'), {}).get('title', 'Session')
            data.update(rows=requests, columns=[('student_name', 'Student'), ('activity_title', 'Session'), ('requested_status', 'Requested'), ('reason', 'Reason'), ('status', 'Status')])
        elif section == 'student-reports':
            reports = _list(request, 'student_reports', order='created_at.desc')
            students = {item['id']: item for item in _list(request, 'students')}
            for item in reports:
                item['student_name'] = students.get(item['student_id'], {}).get('full_name', 'Student')
            data.update(rows=reports, columns=[('student_name', 'Student'), ('title', 'Subject'), ('body', 'Report'), ('status', 'Status'), ('admin_response', 'Response')])
        elif section == 'funds':
            students = _list(request, 'students', order='serial_number.asc')
            payments = _list(request, 'payments')
            paid = {}
            total_collected = Decimal('0')
            for payment in payments:
                amount = Decimal(str(payment['amount']))
                paid[payment['student_id']] = paid.get(payment['student_id'], Decimal('0')) + amount
                total_collected += amount
            total_expected = sum((Decimal(str(student['trip_fee'])) for student in students), Decimal('0'))
            for student in students:
                student['paid_total'] = paid.get(student['id'], Decimal('0'))
                student['balance'] = Decimal(str(student['trip_fee'])) - student['paid_total']
            data.update(rows=students, columns=[('serial_number', '#'), ('full_name', 'Student'), ('trip_fee', 'Fee'), ('paid_total', 'Paid'), ('balance', 'Balance')], summary={'Students': len(students), 'Expected': total_expected, 'Collected': total_collected, 'Balance': total_expected - total_collected})
        elif section == 'photos':
            photos = _list(request, 'photos', order='created_at.desc')
            for photo in photos:
                photo['public_url'] = supabase.storage_public_url('trip-photos', photo['storage_path'])
            data.update(rows=photos, columns=[('caption', 'Caption'), ('storage_path', 'Storage path'), ('created_at', 'Uploaded')])
        elif section == 'notifications':
            data.update(rows=_list(request, 'notifications', order='created_at.desc', limit='50'), columns=[('title', 'Title'), ('message', 'Message'), ('type', 'Type'), ('created_at', 'Created')])
        elif section == 'settings':
            data['rows'] = [_one(request, 'trips'), _one(request, 'trip_settings')]
            data['rows'] = [row for row in data['rows'] if row]
            data['columns'] = [('name', 'Name'), ('start_date', 'Start'), ('end_date', 'End'), ('status', 'Status')]
    except supabase.SupabaseError as error:
        _api_error(request, error)
    return render(request, 'portal/table.html', {'title': title, **data})


@auth.role_required('student')
def student_section(request, section):
    section_titles = {
        'home': 'My Trip', 'schedule': 'Trip Schedule', 'attendance': 'Attendance',
        'attendance-request': 'Attendance Request', 'report': 'Submit Report',
        'photos': 'Trip Photos', 'payments': 'My Payments', 'profile': 'My Profile',
        'notifications': 'Notifications',
    }
    if section not in section_titles:
        return redirect('student-section', section='home')
    data = {'section': section, 'rows': [], 'columns': [], 'summary': {}}
    try:
        student = _student_for_profile(request)
        data['student'] = student
        if section in {'home', 'schedule'}:
            trip = _one(request, 'trips')
            days = _list(request, 'trip_days', order='day_number.asc')
            activities = _list(request, 'activities', order='sort_order.asc')
            for day in days:
                day['activities'] = [activity for activity in activities if activity['trip_day_id'] == day['id']]
            data.update(trip=trip, rows=days, columns=[('day_number', 'Day'), ('trip_date', 'Date'), ('title', 'Itinerary')])
        elif section == 'payments':
            payments = _list(request, 'payments', student_id=f"eq.{student['id']}", order='payment_date.desc') if student else []
            total_paid = sum((Decimal(str(payment['amount'])) for payment in payments), Decimal('0'))
            data.update(rows=payments, columns=[('payment_date', 'Date'), ('amount', 'Amount'), ('payment_method', 'Method'), ('notes', 'Notes')], summary={'Fee': Decimal(str(student['trip_fee'])) if student else Decimal('0'), 'Paid': total_paid, 'Balance': Decimal(str(student['trip_fee'])) - total_paid if student else Decimal('0')})
        elif section == 'attendance':
            sessions = _list(request, 'attendance_sessions', status='eq.active', order='started_at.desc')
            activities = {activity['id']: activity for activity in _list(request, 'activities')}
            records = _list(request, 'attendance_records', student_id=f"eq.{student['id']}") if student else []
            for session in sessions:
                session['activity_title'] = activities.get(session['activity_id'], {}).get('title', 'Activity')
                session['already_marked'] = any(record['session_id'] == session['id'] for record in records)
            data.update(rows=sessions, columns=[('activity_title', 'Activity'), ('started_at', 'Opened'), ('already_marked', 'Marked')])
        elif section == 'attendance-request':
            data['rows'] = _list(request, 'attendance_requests', student_id=f"eq.{student['id']}", order='created_at.desc') if student else []
            data['sessions'] = _list(request, 'attendance_sessions', order='started_at.desc')
            data['activities'] = _list(request, 'activities')
        elif section == 'report':
            data['rows'] = _list(request, 'student_reports', student_id=f"eq.{student['id']}", order='created_at.desc') if student else []
        elif section == 'photos':
            data['rows'] = _list(request, 'photos', order='created_at.desc')
            for photo in data['rows']:
                photo['public_url'] = supabase.storage_public_url('trip-photos', photo['storage_path'])
        elif section == 'notifications':
            data['rows'] = _list(request, 'notifications', user_id=f"eq.{request.profile['id']}", order='created_at.desc', limit='50')
        elif section == 'profile':
            data['profile'] = request.profile
    except (supabase.SupabaseError, TypeError) as error:
        _api_error(request, error)
    return render(request, 'portal/student.html', {'title': section_titles[section], **data})


@auth.role_required('principal')
def principal_section(request, section):
    section_titles = {'home': 'Principal View', 'schedule': 'Schedule', 'activities': 'Activities', 'photos': 'Photos'}
    if section not in section_titles:
        return redirect('principal-section', section='home')
    data = {'section': section, 'rows': [], 'columns': []}
    try:
        if section in {'home', 'schedule'}:
            days = _list(request, 'trip_days', order='day_number.asc')
            activities = _list(request, 'activities', order='sort_order.asc')
            for day in days:
                day['activities'] = [activity for activity in activities if activity['trip_day_id'] == day['id']]
            data.update(trip=_one(request, 'trips'), rows=days, columns=[('day_number', 'Day'), ('trip_date', 'Date'), ('title', 'Itinerary')])
        elif section == 'activities':
            data.update(rows=_list(request, 'activities', order='start_time.asc'), columns=[('title', 'Activity'), ('location', 'Location'), ('start_time', 'Start'), ('status', 'Status')])
        else:
            photos = _list(request, 'photos', order='created_at.desc')
            for photo in photos:
                photo['public_url'] = supabase.storage_public_url('trip-photos', photo['storage_path'])
            data.update(rows=photos, columns=[('caption', 'Caption'), ('storage_path', 'Storage path'), ('created_at', 'Uploaded')])
    except supabase.SupabaseError as error:
        _api_error(request, error)
    return render(request, 'portal/table.html', {'title': section_titles[section], **data})


@auth.role_required('admin')
def student_create(request):
    if request.method != 'POST':
        return render(request, 'portal/form.html', {'title': 'Add Student', 'form_kind': 'student'})
    name = request.POST.get('full_name', '').strip()
    admission = request.POST.get('admission_number', '').strip()
    password = request.POST.get('password', '')
    phone = request.POST.get('phone', '').strip()
    try:
        fee = Decimal(request.POST.get('trip_fee', '6450'))
    except InvalidOperation:
        fee = Decimal('6450')
    if not name or not admission or len(password) < 8:
        messages.error(request, 'Name, admission number, and a password of at least 8 characters are required.')
        return render(request, 'portal/form.html', {'title': 'Add Student', 'form_kind': 'student', 'form_data': request.POST})
    try:
        account = supabase.auth_admin_create(f'{admission}@yaaron.com', password, name)
        profile = supabase.rest_insert('profiles', {'user_id': account['id'], 'full_name': name, 'phone': phone or None, 'role': 'student'}, token=request.supabase_token)
        trip = _one(request, 'trips')
        if not trip:
            raise supabase.SupabaseError('Create a trip before adding students.')
        supabase.rest_insert('students', {
            'profile_id': profile[0]['id'], 'full_name': name, 'admission_number': admission,
            'phone': phone or None, 'class_name': request.POST.get('class_name', '').strip() or None,
            'division': request.POST.get('division', '').strip() or None, 'trip_id': trip['id'], 'trip_fee': float(fee),
        }, token=request.supabase_token)
        messages.success(request, f'Student {admission} was created.')
        return redirect('admin-section', section='students')
    except supabase.SupabaseError as error:
        _api_error(request, error)
    return render(request, 'portal/form.html', {'title': 'Add Student', 'form_kind': 'student', 'form_data': request.POST})


@auth.role_required('admin')
def student_detail(request, student_id):
    try:
        student = _one(request, 'students', id=f'eq.{student_id}')
        if not student:
            return redirect('admin-section', section='students')
        payments = _list(request, 'payments', student_id=f'eq.{student_id}', order='payment_date.desc')
        student['payments'] = payments
        student['total_paid'] = sum((Decimal(str(payment['amount'])) for payment in payments), Decimal('0'))
        student['balance'] = Decimal(str(student['trip_fee'])) - student['total_paid']
        profile = _one(request, 'profiles', id=f"eq.{student['profile_id']}") if student.get('profile_id') else None
        student['auth_user_id'] = profile.get('user_id') if profile else None
        return render(request, 'portal/student_detail.html', {'title': student['full_name'], 'student': student})
    except supabase.SupabaseError as error:
        _api_error(request, error)
        return redirect('admin-section', section='students')


@auth.role_required('admin')
def student_password(request, student_id):
    if request.method != 'POST':
        return HttpResponseNotAllowed(['POST'])
    new_password = request.POST.get('password', '')
    if len(new_password) < 8:
        messages.error(request, 'Password must be at least 8 characters.')
        return redirect('student-detail', student_id=student_id)
    try:
        student = _one(request, 'students', id=f'eq.{student_id}')
        profile = _one(request, 'profiles', id=f"eq.{student['profile_id']}") if student and student.get('profile_id') else None
        if not profile:
            raise supabase.SupabaseError('Student account profile is missing.')
        supabase.auth_admin_update(profile['user_id'], {'password': new_password})
        messages.success(request, 'Student password updated.')
    except supabase.SupabaseError as error:
        _api_error(request, error)
    return redirect('student-detail', student_id=student_id)


@auth.role_required('admin')
def payment_create(request, student_id):
    if request.method != 'POST':
        return HttpResponseNotAllowed(['POST'])
    try:
        amount = Decimal(request.POST.get('amount', '0'))
        if amount <= 0:
            raise InvalidOperation
    except InvalidOperation:
        messages.error(request, 'Enter a valid payment amount.')
        return redirect('student-detail', student_id=student_id)
    try:
        supabase.rest_insert('payments', {
            'student_id': str(student_id), 'amount': float(amount),
            'payment_method': request.POST.get('payment_method', 'cash'),
            'payment_date': request.POST.get('payment_date') or timezone.localdate().isoformat(),
            'notes': request.POST.get('notes', '').strip() or None,
            'recorded_by': request.profile['id'],
        }, token=request.supabase_token)
        messages.success(request, 'Payment recorded.')
    except supabase.SupabaseError as error:
        _api_error(request, error)
    return redirect('student-detail', student_id=student_id)


@auth.role_required('admin')
def student_delete(request, student_id):
    if request.method != 'POST':
        return HttpResponseNotAllowed(['POST'])
    try:
        supabase.rest_delete('students', {'id': f'eq.{student_id}'}, token=request.supabase_token)
        messages.success(request, 'Student and related attendance/payment records were deleted.')
    except supabase.SupabaseError as error:
        _api_error(request, error)
    return redirect('admin-section', section='students')


@auth.role_required('admin')
def payment_delete(request, payment_id):
    if request.method != 'POST':
        return HttpResponseNotAllowed(['POST'])
    try:
        supabase.rest_delete('payments', {'id': f'eq.{payment_id}'}, token=request.supabase_token)
        messages.success(request, 'Payment entry deleted.')
    except supabase.SupabaseError as error:
        _api_error(request, error)
    return redirect('admin-section', section='students')


@auth.role_required('admin')
def photo_upload(request):
    if request.method != 'POST':
        return HttpResponseNotAllowed(['POST'])
    uploaded_files = request.FILES.getlist('photos')
    if not uploaded_files:
        messages.error(request, 'Select at least one image to upload.')
        return redirect('admin-section', section='photos')
    try:
        trip = _one(request, 'trips')
        if not trip:
            raise supabase.SupabaseError('No trip is configured.')
        for upload in uploaded_files:
            if upload.content_type not in {'image/jpeg', 'image/png', 'image/webp'} or upload.size > 10 * 1024 * 1024:
                raise supabase.SupabaseError('Only JPG, PNG, or WEBP images up to 10 MB are supported.')
            safe_name = upload.name.replace('\\', '/').split('/')[-1]
            path = f"{timezone.now():%Y%m%d%H%M%S%f}-{safe_name}"
            supabase.storage_upload('trip-photos', path, upload.read(), upload.content_type, request.supabase_token)
            supabase.rest_insert('photos', {
                'trip_id': trip['id'], 'storage_path': path,
                'caption': request.POST.get('caption', '').strip() or None,
                'uploaded_by': request.profile['id'],
            }, token=request.supabase_token)
        messages.success(request, f'Uploaded {len(uploaded_files)} photo(s).')
    except supabase.SupabaseError as error:
        _api_error(request, error)
    return redirect('admin-section', section='photos')


@auth.role_required('admin')
def activity_status(request, activity_id):
    if request.method != 'POST':
        return HttpResponseNotAllowed(['POST'])
    try:
        supabase.rest_update('activities', {'id': f'eq.{activity_id}'}, {'status': request.POST.get('status')}, token=request.supabase_token)
        messages.success(request, 'Activity status updated.')
    except supabase.SupabaseError as error:
        _api_error(request, error)
    return redirect('admin-section', section='activities')


@auth.role_required('admin')
def attendance_start(request, activity_id):
    if request.method != 'POST':
        return HttpResponseNotAllowed(['POST'])
    try:
        supabase.rest_insert('attendance_sessions', {'activity_id': str(activity_id), 'started_by': request.profile['id']}, token=request.supabase_token)
        messages.success(request, 'Attendance session opened.')
    except supabase.SupabaseError as error:
        _api_error(request, error)
    return redirect('admin-section', section='attendance')


@auth.role_required('admin')
def attendance_close(request, session_id):
    if request.method != 'POST':
        return HttpResponseNotAllowed(['POST'])
    try:
        supabase.rest_update('attendance_sessions', {'id': f'eq.{session_id}'}, {'status': 'closed', 'closed_at': timezone.now().isoformat()}, token=request.supabase_token)
        messages.success(request, 'Attendance session closed.')
    except supabase.SupabaseError as error:
        _api_error(request, error)
    return redirect('admin-section', section='attendance')


@auth.role_required('student')
def attendance_mark(request):
    if request.method != 'POST':
        return HttpResponseNotAllowed(['POST'])
    try:
        student = _student_for_profile(request)
        if not student:
            raise supabase.SupabaseError('No student record is linked to your profile.')
        supabase.rest_insert('attendance_records', {
            'session_id': request.POST.get('session_id'), 'student_id': student['id'], 'status': 'present',
        }, token=request.supabase_token)
        messages.success(request, 'Attendance marked present.')
    except supabase.SupabaseError as error:
        _api_error(request, error)
    return redirect('student-section', section='attendance')


@auth.role_required('student')
def attendance_request_submit(request):
    if request.method != 'POST':
        return HttpResponseNotAllowed(['POST'])
    try:
        student = _student_for_profile(request)
        if not student:
            raise supabase.SupabaseError('No student record is linked to your profile.')
        supabase.rest_insert('attendance_requests', {
            'student_id': student['id'], 'session_id': request.POST.get('session_id'),
            'requested_status': request.POST.get('requested_status'), 'reason': request.POST.get('reason', '').strip(),
        }, token=request.supabase_token)
        messages.success(request, 'Attendance request submitted.')
    except supabase.SupabaseError as error:
        _api_error(request, error)
    return redirect('student-section', section='attendance-request')


@auth.role_required('student')
def student_report_submit(request):
    if request.method != 'POST':
        return HttpResponseNotAllowed(['POST'])
    try:
        student = _student_for_profile(request)
        if not student:
            raise supabase.SupabaseError('No student record is linked to your profile.')
        supabase.rest_insert('student_reports', {
            'student_id': student['id'], 'title': request.POST.get('title', '').strip(), 'body': request.POST.get('body', '').strip(),
        }, token=request.supabase_token)
        messages.success(request, 'Report submitted.')
    except supabase.SupabaseError as error:
        _api_error(request, error)
    return redirect('student-section', section='report')


@auth.role_required('admin')
def attendance_request_review(request, request_id):
    if request.method != 'POST':
        return HttpResponseNotAllowed(['POST'])
    try:
        attendance_request = _one(request, 'attendance_requests', id=f'eq.{request_id}', status='eq.pending')
        if not attendance_request:
            raise supabase.SupabaseError('Request not found or already reviewed.')
        decision = request.POST.get('decision')
        if decision == 'approved':
            supabase.rest_upsert('attendance_records', {
                'session_id': attendance_request['session_id'], 'student_id': attendance_request['student_id'],
                'status': attendance_request['requested_status'], 'marked_at': timezone.now().isoformat(),
            }, conflict='session_id,student_id', token=request.supabase_token)
        if decision not in {'approved', 'rejected'}:
            raise supabase.SupabaseError('Choose approve or reject.')
        supabase.rest_update('attendance_requests', {'id': f'eq.{request_id}', 'status': 'eq.pending'}, {
            'status': decision, 'review_note': request.POST.get('review_note', '').strip() or None,
            'reviewed_by': request.profile['id'], 'reviewed_at': timezone.now().isoformat(),
        }, token=request.supabase_token)
        messages.success(request, 'Attendance request reviewed.')
    except supabase.SupabaseError as error:
        _api_error(request, error)
    return redirect('admin-section', section='attendance-requests')


@auth.role_required('admin')
def student_report_review(request, report_id):
    if request.method != 'POST':
        return HttpResponseNotAllowed(['POST'])
    try:
        supabase.rest_update('student_reports', {'id': f'eq.{report_id}'}, {
            'admin_response': request.POST.get('admin_response', '').strip() or None,
            'status': 'reviewed', 'reviewed_by': request.profile['id'], 'reviewed_at': timezone.now().isoformat(),
        }, token=request.supabase_token)
        messages.success(request, 'Student report response saved.')
    except supabase.SupabaseError as error:
        _api_error(request, error)
    return redirect('admin-section', section='student-reports')


@auth.role_required('admin')
def student_import(request):
    if request.method == 'POST':
        upload = request.FILES.get('workbook')
        if not upload:
            messages.error(request, 'Choose an Excel workbook to import.')
            return redirect('student-import')
        try:
            workbook = load_workbook(upload, read_only=True, data_only=True)
            sheet = next((workbook[name] for name in workbook.sheetnames if 'trip' in name.lower() or 'fund' in name.lower()), workbook.active)
            rows = list(sheet.iter_rows(values_only=True))
            header_index = next((index for index, row in enumerate(rows) if _is_student_header(row)), None)
            if header_index is None:
                raise ValueError('Could not find a student header row in the workbook.')
            headers = [_normalize_header(value) for value in rows[header_index]]
            columns = {
                'name': _find_column(headers, ('name', 'student', 'fullname')),
                'serial': _find_column(headers, ('serial', 'sno', 'slno', 'rollno')),
                'admission': _find_column(headers, ('admission', 'admno')),
                'class_name': _find_column(headers, ('class', 'grade', 'section')),
                'fee': _find_column(headers, ('total', 'fee', 'amount')),
                'paid': _find_column(headers, ('paid', 'received', 'collected', 'advance')),
            }
            source = []
            for row in rows[header_index + 1:]:
                name = _cell(row, columns['name'])
                if not name or name.lower() == 'total':
                    continue
                serial_value = _cell(row, columns['serial'])
                serial_number = int(float(serial_value)) if serial_value else len(source) + 1
                admission_number = _cell(row, columns['admission']) or f'IVS3-{serial_number:03d}'
                try:
                    fee = Decimal(_cell(row, columns['fee']) or '6450')
                    paid = Decimal(_cell(row, columns['paid']) or '0')
                except InvalidOperation as error:
                    raise ValueError(f'Invalid fee/payment amount for serial {serial_number}.') from error
                if fee < 0 or paid < 0 or paid > fee:
                    raise ValueError(f'Invalid fee/payment values for serial {serial_number}.')
                source.append({
                    'serial': serial_number, 'name': name, 'admission': admission_number,
                    'class_name': _cell(row, columns['class_name']) or None, 'fee': fee, 'paid': paid,
                })
            workbook.close()
            if not source:
                raise ValueError('The workbook contains no student rows.')
            if len({student['admission'] for student in source}) != len(source):
                raise ValueError('The workbook contains duplicate admission numbers.')

            trips = supabase.rest_list('trips', token=request.supabase_token)
            if len(trips) != 1:
                raise ValueError('Expected one configured trip before importing students.')
            existing_students = supabase.rest_list('students', {'limit': '1000'}, token=request.supabase_token)
            existing_payments = supabase.rest_list('payments', {'limit': '1000'}, token=request.supabase_token)
            existing_auth = {}
            for page in range(1, 20):
                result = supabase.auth_admin_list_users(page=page, per_page=1000)
                accounts = result.get('users', [])
                for account in accounts:
                    if account.get('email'):
                        existing_auth[account['email'].lower()] = account
                if len(accounts) < 1000:
                    break
            payment_totals = {}
            for payment in existing_payments:
                payment_totals[payment['student_id']] = payment_totals.get(payment['student_id'], Decimal('0')) + Decimal(str(payment['amount']))
            students_by_admission = {row.get('admission_number'): row for row in existing_students if row.get('admission_number')}
            students_by_profile = {row.get('profile_id'): row for row in existing_students if row.get('profile_id')}

            imported = 0
            failures = []
            for student in source:
                login_email = f"{student['admission']}@yaaron.com"
                try:
                    account = existing_auth.get(login_email.lower())
                    if account:
                        auth_user = supabase.auth_admin_update(account['id'], {
                            'password': '12345678', 'user_metadata': {'full_name': student['name']},
                        })
                    else:
                        auth_user = supabase.auth_admin_create(login_email, '12345678', student['name'])
                    user_id = auth_user['id']

                    linked_profiles = supabase.rest_list('profiles', {'user_id': f'eq.{user_id}'}, token=request.supabase_token)
                    if linked_profiles:
                        profile = linked_profiles[0]
                        profile_id = profile['id']
                        if profile.get('role') != 'student' or profile.get('full_name') != student['name']:
                            supabase.rest_update('profiles', {'id': f'eq.{profile_id}'}, {'role': 'student', 'full_name': student['name']}, token=request.supabase_token)
                    else:
                        profile_rows = supabase.rest_insert('profiles', {'user_id': user_id, 'role': 'student', 'full_name': student['name']}, token=request.supabase_token)
                        profile_id = profile_rows[0]['id']

                    existing_student = students_by_admission.get(student['admission']) or students_by_profile.get(profile_id)
                    student_values = {
                        'profile_id': profile_id, 'serial_number': student['serial'], 'full_name': student['name'],
                        'admission_number': student['admission'], 'class_name': student['class_name'],
                        'trip_id': trips[0]['id'], 'trip_fee': float(student['fee']),
                    }
                    if existing_student:
                        supabase.rest_update('students', {'id': f"eq.{existing_student['id']}"}, student_values, token=request.supabase_token)
                        student_id = existing_student['id']
                    else:
                        new_student = supabase.rest_insert('students', student_values, token=request.supabase_token)[0]
                        student_id = new_student['id']
                        students_by_admission[student['admission']] = new_student
                    if student['paid'] > 0 and payment_totals.get(student_id, Decimal('0')) == 0:
                        supabase.rest_insert('payments', {
                            'student_id': student_id, 'amount': float(student['paid']), 'payment_method': 'cash',
                            'notes': 'Imported from spreadsheet', 'recorded_by': request.profile['id'],
                        }, token=request.supabase_token)
                    imported += 1
                except (supabase.SupabaseError, KeyError, IndexError) as error:
                    failures.append(f"{student['admission']}: {error}")

            if failures:
                messages.warning(request, f'Synced {imported} of {len(source)} students. ' + ' '.join(failures[:4]))
            else:
                messages.success(request, f'Synced {imported} students and reset their temporary passwords to 12345678.')
            return redirect('admin-section', section='students')
        except (supabase.SupabaseError, ValueError, InvalidOperation, TypeError) as error:
            messages.error(request, str(error))
    return render(request, 'portal/import.html', {'title': 'Import Students'})


def _is_student_header(row):
    headers = [_normalize_header(value) for value in row]
    return (
        any(any(token in header for token in ('name', 'student', 'fullname')) for header in headers)
        and any(any(token in header for token in ('total', 'fee', 'amount')) for header in headers)
        and any(any(token in header for token in ('paid', 'received', 'collected')) for header in headers)
    )